import io
import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_pdf_report(franchise_dict, timeline_items, financial_summary):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=12
    )
    story.append(Paragraph(f"Franchise Profile Report: {franchise_dict.get('name', 'N/A')}", title_style))
    story.append(Paragraph(f"Code: {franchise_dict.get('code', 'N/A')} | Owner: {franchise_dict.get('owner_name', 'N/A')} | City: {franchise_dict.get('city', 'N/A')}", styles['Normal']))
    story.append(Paragraph(f"Report Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Italic']))
    story.append(Spacer(1, 15))

    # Financial Summary Table
    story.append(Paragraph("<b>Financial Summary</b>", styles['Heading2']))
    fin_data = [
        ["Metric", "Value"],
        ["Total Purchases", f"Rs. {financial_summary.get('total_purchase', 0):,.2f}"],
        ["Total GR Return", f"Rs. {financial_summary.get('total_gr', 0):,.2f}"],
        ["GR Percentage (GR%)", f"{financial_summary.get('gr_percent', 0):.2f}%"],
        ["Net Purchase", f"Rs. {financial_summary.get('net_purchase', 0):,.2f}"],
        ["Total Payment Received", f"Rs. {financial_summary.get('total_received', 0):,.2f}"],
        ["Outstanding Amount", f"Rs. {financial_summary.get('outstanding', 0):,.2f}"],
        ["Company Support Investment", f"Rs. {financial_summary.get('total_support', 0):,.2f}"]
    ]
    fin_table = Table(fin_data, colWidths=[240, 240])
    fin_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (1,0), colors.HexColor('#3B82F6')),
        ('TEXTCOLOR', (0,0), (1,0), colors.white),
        ('FONTNAME', (0,0), (1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(fin_table)
    story.append(Spacer(1, 15))

    # Timeline / Activity History
    story.append(Paragraph("<b>Chronological Activity Timeline</b>", styles['Heading2']))
    timeline_data = [["Date / Time", "Stage / Module", "Person", "Action / Details"]]
    for item in timeline_items:
        timeline_data.append([
            str(item.get('timestamp', item.get('created_at', ''))),
            str(item.get('stage_name', item.get('module', ''))),
            str(item.get('person', item.get('performed_by', ''))),
            str(item.get('remarks', item.get('details', '')))[:60]
        ])
    if len(timeline_data) == 1:
        timeline_data.append(["-", "No recorded activities yet", "-", "-"])

    timeline_table = Table(timeline_data, colWidths=[110, 110, 100, 160])
    timeline_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F1F5F9')])
    ]))
    story.append(timeline_table)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_report(franchise_list, stage_entries):
    wb = openpyxl.Workbook()
    
    # Sheet 1: Franchises Overview
    ws1 = wb.active
    ws1.title = "Franchises Overview"
    headers1 = ["ID", "Code", "Name", "Owner Name", "Mobile", "City", "Assigned Person", "Status", "Plan", "Agreed Amount"]
    ws1.append(headers1)
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_num, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for f in franchise_list:
        ws1.append([
            f.get('id'), f.get('code'), f.get('name'), f.get('owner_name'), f.get('owner_mobile'),
            f.get('city'), f.get('assigned_person'), f.get('status'), f.get('plan_name'), f.get('agreed_amount')
        ])

    # Sheet 2: Stage Entries
    ws2 = wb.create_sheet(title="Stage & Audit Log")
    headers2 = ["Date", "Franchise Code", "Stage", "Person", "Action", "Remarks"]
    ws2.append(headers2)
    for col_num, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for entry in stage_entries:
        ws2.append([
            entry.get('timestamp', entry.get('created_at', '')),
            entry.get('franchise_code', ''),
            entry.get('stage_name', ''),
            entry.get('performed_by', entry.get('person', '')),
            entry.get('action', ''),
            entry.get('remarks', '')
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_word_report(franchise_dict, financial_summary):
    # HTML formatted document export for Word (.doc format)
    html_content = f"""
    <html xmlns:o='urn:schemas-microsoft-com:office:office' xmlns:w='urn:schemas-microsoft-com:office:word' xmlns='http://www.w3.org/TR/REC-html40'>
    <head><title>Franchise Profile Report</title>
    <style>
        body {{ font-family: Calibri, sans-serif; margin: 20px; }}
        h1 {{ color: #1E293B; border-bottom: 2px solid #3B82F6; padding-bottom: 5px; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 15px; }}
        th {{ background-color: #3B82F6; color: white; border: 1px solid #ddd; padding: 8px; text-align: left; }}
        td {{ border: 1px solid #ddd; padding: 8px; }}
        tr:nth-child(even) {{ background-color: #f2f2f2; }}
    </style>
    </head>
    <body>
        <h1>Franchise Executive Summary</h1>
        <p><b>Franchise Name:</b> {franchise_dict.get('name', 'N/A')}</p>
        <p><b>Code:</b> {franchise_dict.get('code', 'N/A')}</p>
        <p><b>Owner:</b> {franchise_dict.get('owner_name', 'N/A')} ({franchise_dict.get('owner_mobile', 'N/A')})</p>
        <p><b>City:</b> {franchise_dict.get('city', 'N/A')}</p>
        <p><b>Assigned Manager:</b> {franchise_dict.get('assigned_person', 'N/A')}</p>
        <p><b>Status:</b> {franchise_dict.get('status', 'N/A')}</p>

        <h2>Financial Metrics</h2>
        <table>
            <tr><th>Metric</th><th>Amount / Value</th></tr>
            <tr><td>Total Purchases</td><td>Rs. {financial_summary.get('total_purchase', 0):,.2f}</td></tr>
            <tr><td>Total GR Returns</td><td>Rs. {financial_summary.get('total_gr', 0):,.2f}</td></tr>
            <tr><td>GR Percentage (GR%)</td><td>{financial_summary.get('gr_percent', 0):.2f}%</td></tr>
            <tr><td>Net Purchase</td><td>Rs. {financial_summary.get('net_purchase', 0):,.2f}</td></tr>
            <tr><td>Payment Received</td><td>Rs. {financial_summary.get('total_received', 0):,.2f}</td></tr>
            <tr><td>Outstanding Amount</td><td>Rs. {financial_summary.get('outstanding', 0):,.2f}</td></tr>
            <tr><td>Total Company Support Investment</td><td>Rs. {financial_summary.get('total_support', 0):,.2f}</td></tr>
        </table>
    </body>
    </html>
    """
    return html_content.encode('utf-8')
